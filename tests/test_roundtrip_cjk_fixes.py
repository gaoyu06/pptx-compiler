#!/usr/bin/env python3
"""Regression tests for round-trip editing and CJK text handling.

- An edited round-trip slide whose first object is a full-canvas source
  picture must still export (background promotion used to swallow it).
- A font family named in CJK script must fill the East Asian slot.
- Ambiguous-width CJK punctuation must be sized full-width.
"""

from __future__ import annotations

import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Emu


SCRIPTS_DIR = Path(__file__).resolve().parents[1]
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

from svg_to_pptx.drawingml.utils import (  # noqa: E402
    estimate_text_width,
    parse_font_family,
)


def _run(script: str, *args: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, str(SCRIPTS_DIR / script), *args],
        capture_output=True,
        text=True,
    )


class RoundtripFullCanvasPictureTests(unittest.TestCase):
    def test_edited_slide_keeps_full_canvas_source_picture(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            photo = tmp_path / "photo.png"
            Image.new("RGB", (64, 36), "#C53826").save(photo)

            prs = Presentation()
            prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            slide.shapes.add_picture(
                str(photo), 0, 0, prs.slide_width, prs.slide_height
            )
            box = slide.shapes.add_textbox(Emu(914400), Emu(914400), Emu(3657600), Emu(914400))
            box.text_frame.text = "original"
            source = tmp_path / "deck.pptx"
            prs.save(source)

            workspace = tmp_path / "ws"
            imported = _run(
                "pptx_to_svg.py", str(source), "-o", str(workspace), "--roundtrip"
            )
            self.assertEqual(imported.returncode, 0, imported.stderr)

            page = workspace / "authoring-svg-flat" / "slide_01.svg"
            page.write_text(
                page.read_text(encoding="utf-8").replace("original", "edited"),
                encoding="utf-8",
            )

            output = tmp_path / "out.pptx"
            exported = _run(
                "svg_to_pptx.py", str(workspace), "--roundtrip", "-o", str(output)
            )
            self.assertEqual(exported.returncode, 0, exported.stdout + exported.stderr)

            with zipfile.ZipFile(output) as package:
                slide_xml = package.read("ppt/slides/slide1.xml").decode("utf-8")
            self.assertIn("<p:pic", slide_xml)
            self.assertIn("edited", slide_xml)


class CjkFontFamilyTests(unittest.TestCase):
    def test_cjk_named_family_fills_east_asian_slot(self) -> None:
        fonts = parse_font_family('"思源黑体 CN Medium", sans-serif')
        self.assertEqual(fonts["ea"], "思源黑体 CN Medium")
        self.assertEqual(fonts["latin"], "思源黑体 CN Medium")

    def test_latin_family_before_cjk_named_family(self) -> None:
        fonts = parse_font_family('Montserrat, "思源黑体 CN Medium", sans-serif')
        self.assertEqual(fonts["latin"], "Montserrat")
        self.assertEqual(fonts["ea"], "思源黑体 CN Medium")


class CjkPunctuationWidthTests(unittest.TestCase):
    def test_ambiguous_cjk_punctuation_is_full_width(self) -> None:
        for ch in "“”‘’·—…":
            with self.subTest(ch=ch):
                self.assertEqual(estimate_text_width(ch, 20), 20)


if __name__ == "__main__":
    unittest.main()
